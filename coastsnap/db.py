"""Site metadata from ``CoastSnapDB.xlsx`` (port of CSPreadSiteDB.m).

Each site sheet is a key/value layout: labels in column A, values in column B.
"""

from __future__ import annotations



def _rows(path, site):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[site]
        return [(r[0], r[1] if len(r) > 1 else None)
                for r in ws.iter_rows(min_col=1, max_col=2, values_only=True)]
    finally:
        wb.close()


def _num_list(v):
    """MATLAB ``str2num`` on strings such as ``"[1 2 3]"`` or ``"1,2,3"``."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return [v]
    s = str(v).strip().strip("[]").replace(",", " ").replace(";", " ")
    return [float(t) if "." in t else int(t) for t in s.split()]


def read_site_db(path: str, site: str) -> dict:
    """Read one site sheet into the same nested structure as CSPreadSiteDB."""
    rows = _rows(path, site)
    labels = [str(a).strip() if a is not None else None for a, _ in rows]

    def val(label, offset=0):
        try:
            i = labels.index(label)
        except ValueError:
            return None
        return rows[i + offset][1]

    out = {
        "origin": {"eastings": val("Station Data", 1), "northings": val("Station Data", 2),
                   "z": val("Station Data", 3)},
        "UTMzone": val("UTM Zone"),
        "timezone": {
            "name": val("Default Timezone"),
            "gmt_offset": val("Default Timezone Offset From GMT"),
            "alternative": {"name": val("Alternative Timezone"),
                            "gmt_offset": val("Alternative Timezone Offset From GMT")},
        },
        "rect": {
            "xlim": [val("Xlimit left"), val("Xlimit right")],
            "ylim": [val("Ylimit lower"), val("Ylimit upper")],
            "res": val("Resolution"),
            "initial": {"azimuth": val("Initial Azimuth Estimate"),
                        "tilt": val("Initial Tilt Estimate"),
                        "roll": val("Initial Roll Estimate")},
            "tidal_offset": val("Tidal offset"),
            "FOVlims": [val("Min FOV"), val("Max FOV")],
            "accuracylim": val("Acceptable Accuracy"),
        },
        "tide": {"file": val("Tide file")},
        "sl_settings": {
            "transect_file": val("Transect file"),
            "transect_averaging_region": _num_list(val("Transect averaging region")),
            "beach_slope": val("Characteristic beach slope"),
        },
        "gcp": [
            {"name": str(rows[i][1]), "eastings": rows[i + 1][1],
             "northings": rows[i + 2][1], "z": rows[i + 3][1]}
            for i, lab in enumerate(labels) if lab == "GCP name"
        ],
        "gcp_combo": _num_list(val("GCP combo")),
    }
    if None in out["rect"]["FOVlims"]:
        out["rect"]["FOVlims"] = None
    return out
